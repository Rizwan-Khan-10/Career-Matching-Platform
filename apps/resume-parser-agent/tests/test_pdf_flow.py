"""PDF delivery flow with Cloudinary/DB/Redis faked. Uses REAL conversion (LibreOffice if installed, else fallback)."""
import io
import sys
import types
import unittest

# stub the redis package (no server / package needed)
if "redis" not in sys.modules:
    m = types.ModuleType("redis")
    m.from_url = lambda *a, **k: None
    m.exceptions = types.SimpleNamespace(ResponseError=Exception)
    sys.modules["redis"] = m

from app.services import pdf_storage                      # noqa: E402
from app.workers import pdf_consumer as pc                # noqa: E402


def make_docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("Asha Verma", 0)
    d.add_paragraph("Backend Developer")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_pdf() -> bytes:
    from reportlab.pdfgen import canvas
    b = io.BytesIO()
    c = canvas.Canvas(b)
    c.drawString(72, 700, "hello")
    c.save()
    return b.getvalue()


class Base(unittest.TestCase):
    def setUp(self):
        self.published, self.uploads, self.saved = [], {}, {}
        self.files = {}
        self.resumes = {}
        pc.publish = lambda stream, data: self.published.append((stream, data))
        pdf_storage.download_original = lambda pid: self.files[pid]
        pdf_storage.upload_pdf = lambda pid, b: self.uploads.__setitem__(pid, b)
        pc._load_resume = lambda rid, aid: self.resumes.get(rid)
        pc._save_pdf_public_id = lambda rid, pid: self.saved.__setitem__(rid, pid)
        self.event = {"matchId": "m1", "applicantId": "a1", "jobRoleId": "r1", "resumeId": "res1"}

    def add_resume(self, public_id, content, pdf_public_id=None):
        self.resumes["res1"] = {"id": "res1", "applicantId": "a1", "cloudinaryPublicId": public_id, "pdfPublicId": pdf_public_id}
        self.files[public_id] = content


class FlowTest(Base):
    def test_docx_is_converted_uploaded_and_announced(self):
        self.add_resume("resumes/a1/1700000000000-Asha CV.docx", make_docx())
        pc.handle(self.event)
        self.assertEqual(list(self.uploads), ["resumes/a1/pdf/res1.pdf"])
        self.assertTrue(self.uploads["resumes/a1/pdf/res1.pdf"].startswith(b"%PDF"))
        self.assertEqual(self.saved, {"res1": "resumes/a1/pdf/res1.pdf"})
        stream, payload = self.published[-1]
        self.assertEqual(stream, "resume.pdf.ready")
        self.assertEqual(payload["pdfPublicId"], "resumes/a1/pdf/res1.pdf")
        self.assertEqual(payload["fileName"], "Asha_CV.pdf")
        self.assertEqual((payload["matchId"], payload["jobRoleId"], payload["applicantId"]), ("m1", "r1", "a1"))

    def test_pdf_is_reused_without_upload(self):
        pid = "resumes/a1/1700000000000-cv.pdf"
        self.add_resume(pid, make_pdf())
        pc.handle(self.event)
        self.assertEqual(self.uploads, {})
        self.assertEqual(self.saved, {"res1": pid})
        self.assertEqual(self.published[-1][1]["pdfPublicId"], pid)

    def test_existing_pdf_is_not_converted_again(self):
        self.add_resume("resumes/a1/1700000000000-cv.docx", b"would fail if downloaded", pdf_public_id="resumes/a1/pdf/res1.pdf")
        pdf_storage.download_original = lambda pid: self.fail("must not download again")
        pc.handle(self.event)
        self.assertEqual(self.published[-1][0], "resume.pdf.ready")

    def test_corrupt_file_reports_failure_not_retry(self):
        self.add_resume("resumes/a1/1700000000000-cv.bin", b"\x00\x01\x02garbage\xff" * 20)
        pc.handle(self.event)                       # must NOT raise (a retry could never fix it)
        self.assertEqual(self.published[-1][0], "resume.pdf.failed")
        self.assertEqual(self.uploads, {})

    def test_missing_resume_reports_failure(self):
        pc.handle(self.event)
        self.assertEqual(self.published[-1][0], "resume.pdf.failed")

    def test_malformed_event_raises(self):
        with self.assertRaises(ValueError):
            pc.handle({"applicantId": "a1"})

    def test_infra_error_propagates_for_retry(self):
        self.add_resume("resumes/a1/1700000000000-cv.docx", make_docx())
        pdf_storage.upload_pdf = lambda pid, b: (_ for _ in ()).throw(RuntimeError("cloudinary down"))
        with self.assertRaises(RuntimeError):
            pc.handle(self.event)
        self.assertEqual(self.saved, {})            # nothing half-saved
        self.assertEqual(self.published, [])

    def test_on_dead_tells_results_service(self):
        pc.on_dead(self.event, RuntimeError("boom"))
        self.assertEqual(self.published[-1][0], "resume.pdf.failed")


class HelpersTest(unittest.TestCase):
    def test_names_and_extensions(self):
        self.assertEqual(pdf_storage.original_extension("resumes/a/170-My CV.DOCX"), "docx")
        self.assertEqual(pdf_storage.original_extension("resumes/a/170-resume"), "pdf")
        self.assertEqual(pdf_storage.friendly_pdf_name("resumes/a/1700000000000-My CV (final).docx"), "My_CV_final.pdf")
        self.assertEqual(pdf_storage.friendly_pdf_name("resumes/a/1700000000000-.pdf"), "Resume.pdf")


if __name__ == "__main__":
    unittest.main()